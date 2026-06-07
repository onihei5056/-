"""
不動産管理システム — コアロジック (DB・Excel)
"""

import sqlite3
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = os.path.join(os.path.dirname(__file__), "realestate.db")

# ─── カテゴリ定数 ──────────────────────────────────────────────────────────────

SALE_TYPES    = ["売買", "賃貸（仲介）", "賃貸（管理）", "その他"]
PROP_TYPES    = ["マンション", "一戸建て", "土地", "事務所・店舗", "駐車場", "その他"]
STATUSES      = ["販売中", "成約済", "取下げ", "商談中"]
EXPENSE_CATS  = [
    "家賃・地代", "水道光熱費", "通信費", "広告宣伝費",
    "人件費", "交通費", "消耗品費", "接待交際費",
    "保険料", "修繕費", "顧問料", "その他"
]

# ─── Database ──────────────────────────────────────────────────────────────────

def get_conn():
    return sqlite3.connect(DB_PATH)

def init_db():
    with get_conn() as c:
        c.execute("""
            CREATE TABLE IF NOT EXISTS sales (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                date          TEXT NOT NULL,
                property_id   TEXT,
                property_name TEXT NOT NULL,
                type          TEXT NOT NULL,
                price         INTEGER NOT NULL,
                fee           INTEGER NOT NULL,
                staff         TEXT,
                memo          TEXT,
                created_at    TEXT DEFAULT (datetime('now','localtime'))
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS inventory (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                property_no   TEXT,
                property_name TEXT NOT NULL,
                type          TEXT NOT NULL,
                address       TEXT,
                price         INTEGER,
                area          REAL,
                rooms         TEXT,
                status        TEXT NOT NULL DEFAULT '販売中',
                listed_date   TEXT,
                memo          TEXT,
                created_at    TEXT DEFAULT (datetime('now','localtime'))
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS expenses (
                id         INTEGER PRIMARY KEY AUTOINCREMENT,
                date       TEXT NOT NULL,
                category   TEXT NOT NULL,
                amount     INTEGER NOT NULL,
                vendor     TEXT,
                memo       TEXT,
                created_at TEXT DEFAULT (datetime('now','localtime'))
            )
        """)

# ─── ユーティリティ ────────────────────────────────────────────────────────────

def jpy(n):
    try:
        return f"¥{int(n):,}"
    except Exception:
        return "¥0"

def today_str():
    return date.today().strftime("%Y-%m-%d")

# ─── Excel スタイル定数 ────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F497D")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="メイリオ")
TITLE_FONT  = Font(bold=True, size=14, name="メイリオ")
CELL_FONT   = Font(name="メイリオ", size=10)
TOTAL_FILL  = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FONT  = Font(bold=True, name="メイリオ", size=10)
ALT_FILL    = PatternFill("solid", fgColor="EEF2F8")

_thin   = Side(style="thin", color="BBBBBB")
BORDER  = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

def _style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = BORDER

def _style_row(ws, row, cols, alt=False):
    for c in range(1, cols + 1):
        cell        = ws.cell(row=row, column=c)
        cell.font   = CELL_FONT
        cell.border = BORDER
        if alt:
            cell.fill = ALT_FILL

def _autofit(ws, extra=2):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(
            (len(str(cell.value)) for cell in col if cell.value),
            default=0
        )
        ws.column_dimensions[col_letter].width = min(max_len * 1.8 + extra, 50)

# ─── Excel エクスポート ────────────────────────────────────────────────────────

def export_excel(filepath: str, year_month: str | None = None):
    """
    4シート構成のExcelを生成して filepath に保存する。
    year_month が指定された場合は売上・経費を対象月に絞る（例: '2026-05'）。
    """
    wb = Workbook()
    wb.remove(wb.active)

    with get_conn() as c:
        _write_sales_sheet(wb, c, year_month)
        _write_inventory_sheet(wb, c)
        _write_expense_sheet(wb, c, year_month)
        _write_summary_sheet(wb, c)

    wb.save(filepath)
    return filepath


def _write_sales_sheet(wb, c, year_month):
    ws = wb.create_sheet("売上一覧")
    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 28

    headers = ["日付", "物件名", "種別", "成約金額(円)", "仲介手数料(円)", "担当者", "備考"]
    ws.append([f"不動産仲介 売上一覧{' (' + year_month + ')' if year_month else ''}"])
    ws["A1"].font = TITLE_FONT
    ws.append(headers)
    _style_header(ws, 2, len(headers))

    query  = "SELECT date, property_name, type, price, fee, staff, memo FROM sales"
    params = ()
    if year_month:
        query += " WHERE date LIKE ? ORDER BY date"
        params = (f"{year_month}%",)
    else:
        query += " ORDER BY date"

    rows = c.execute(query, params).fetchall()
    total_price = total_fee = 0
    for i, row in enumerate(rows):
        ws.append(list(row))
        r = i + 3
        _style_row(ws, r, len(headers), alt=(i % 2 == 1))
        for col in [4, 5]:
            cell = ws.cell(r, col)
            cell.number_format = '#,##0'
            cell.alignment     = Alignment(horizontal="right")
        total_price += row[3] or 0
        total_fee   += row[4] or 0

    tr = ws.max_row + 1
    ws.cell(tr, 1, "合計")
    ws.cell(tr, 4, total_price).number_format = '#,##0'
    ws.cell(tr, 5, total_fee).number_format   = '#,##0'
    for col in range(1, len(headers) + 1):
        cell        = ws.cell(tr, col)
        cell.fill   = TOTAL_FILL
        cell.font   = TOTAL_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="right" if col in [4, 5] else "left")

    _autofit(ws)


def _write_inventory_sheet(wb, c):
    ws = wb.create_sheet("物件在庫")
    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 28

    headers = ["物件番号", "物件名", "種別", "所在地", "価格(円)", "面積(㎡)", "間取り", "状態", "掲載日", "備考"]
    ws.append(["物件在庫一覧"])
    ws["A1"].font = TITLE_FONT
    ws.append(headers)
    _style_header(ws, 2, len(headers))

    rows = c.execute(
        "SELECT property_no, property_name, type, address, price, area, rooms, status, listed_date, memo"
        " FROM inventory ORDER BY status, listed_date"
    ).fetchall()
    for i, row in enumerate(rows):
        ws.append(list(row))
        r = i + 3
        _style_row(ws, r, len(headers), alt=(i % 2 == 1))
        cell = ws.cell(r, 5)
        cell.number_format = '#,##0'
        cell.alignment     = Alignment(horizontal="right")

    _autofit(ws)


def _write_expense_sheet(wb, c, year_month):
    ws = wb.create_sheet("経費一覧")
    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 28

    headers = ["日付", "費目", "金額(円)", "支払先", "備考"]
    ws.append([f"店舗運営経費一覧{' (' + year_month + ')' if year_month else ''}"])
    ws["A1"].font = TITLE_FONT
    ws.append(headers)
    _style_header(ws, 2, len(headers))

    query  = "SELECT date, category, amount, vendor, memo FROM expenses"
    params = ()
    if year_month:
        query += " WHERE date LIKE ? ORDER BY date"
        params = (f"{year_month}%",)
    else:
        query += " ORDER BY date"

    rows      = c.execute(query, params).fetchall()
    total_exp = 0
    for i, row in enumerate(rows):
        ws.append(list(row))
        r = i + 3
        _style_row(ws, r, len(headers), alt=(i % 2 == 1))
        cell = ws.cell(r, 3)
        cell.number_format = '#,##0'
        cell.alignment     = Alignment(horizontal="right")
        total_exp += row[2] or 0

    tr = ws.max_row + 1
    ws.cell(tr, 1, "合計")
    ws.cell(tr, 3, total_exp).number_format = '#,##0'
    for col in range(1, len(headers) + 1):
        cell        = ws.cell(tr, col)
        cell.fill   = TOTAL_FILL
        cell.font   = TOTAL_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="right" if col == 3 else "left")

    _autofit(ws)


def _write_summary_sheet(wb, c):
    ws = wb.create_sheet("月次サマリー")
    ws.row_dimensions[1].height = 28

    ws.append(["月次損益サマリー"])
    ws["A1"].font = TITLE_FONT
    ws.append(["月", "売上件数", "仲介手数料合計(円)", "経費合計(円)", "利益(円)"])
    _style_header(ws, 2, 5)

    months = c.execute(
        "SELECT DISTINCT substr(date,1,7)"
        " FROM (SELECT date FROM sales UNION SELECT date FROM expenses)"
        " ORDER BY 1"
    ).fetchall()

    for i, (m,) in enumerate(months):
        cnt, fee_sum = c.execute(
            "SELECT COUNT(*), COALESCE(SUM(fee),0) FROM sales WHERE date LIKE ?",
            (f"{m}%",)
        ).fetchone()
        exp_sum = c.execute(
            "SELECT COALESCE(SUM(amount),0) FROM expenses WHERE date LIKE ?",
            (f"{m}%",)
        ).fetchone()[0]
        profit = fee_sum - exp_sum
        ws.append([m, cnt, fee_sum, exp_sum, profit])
        r = i + 3
        _style_row(ws, r, 5, alt=(i % 2 == 1))
        for col in [3, 4, 5]:
            cell = ws.cell(r, col)
            cell.number_format = '#,##0'
            cell.alignment     = Alignment(horizontal="right")
        if profit < 0:
            ws.cell(r, 5).font = Font(color="FF0000", bold=True, name="メイリオ", size=10)

    _autofit(ws)
